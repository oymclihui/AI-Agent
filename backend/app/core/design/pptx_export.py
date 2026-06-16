"""Generate styled PPTX decks aligned with AI-AGENT / MarkdownView light theme."""
from __future__ import annotations

import json
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


@dataclass(frozen=True)
class PptxTheme:
    bg_dark: RGBColor = RGBColor(0x1E, 0x3A, 0x5F)
    bg_light: RGBColor = RGBColor(0xF8, 0xFB, 0xFF)
    accent: RGBColor = RGBColor(0x3B, 0x82, 0xF6)
    accent_soft: RGBColor = RGBColor(0xDB, 0xEA, 0xFE)
    title: RGBColor = RGBColor(0x1E, 0x3A, 0x5F)
    body: RGBColor = RGBColor(0x33, 0x41, 0x55)
    muted: RGBColor = RGBColor(0x5B, 0x7A, 0x9A)
    on_dark: RGBColor = RGBColor(0xFF, 0xFF, 0xFF)
    on_dark_sub: RGBColor = RGBColor(0xBF, 0xDB, 0xFE)
    font_cjk: str = "Microsoft YaHei UI"
    font_latin: str = "Calibri"


THEME = PptxTheme()
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def _parse_slides(slides_json: str) -> list[dict[str, Any]]:
    data = json.loads(slides_json)
    if isinstance(data, dict) and "slides" in data:
        data = data["slides"]
    if not isinstance(data, list) or not data:
        raise ValueError("slides_json must be a non-empty JSON array of slide objects")
    return data


def _blank_layout(prs: Presentation):
    # Prefer blank layout for full theme control
    for idx in (6, 5, len(prs.slide_layouts) - 1):
        if idx < len(prs.slide_layouts):
            return prs.slide_layouts[idx]
    return prs.slide_layouts[0]


def _set_slide_bg(slide, color: RGBColor) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def _set_run_font(run, *, size_pt: float, color: RGBColor, bold: bool = False) -> None:
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = THEME.font_latin
    try:
        run.font._element.rPr.rFonts.set("{http://schemas.openxmlformats.org/drawingml/2006/main}eastAsia", THEME.font_cjk)
    except Exception:
        pass


def _add_textbox(
    slide,
    left,
    top,
    width,
    height,
    text: str,
    *,
    size_pt: float = 18,
    color: RGBColor | None = None,
    bold: bool = False,
    align=PP_ALIGN.LEFT,
) -> None:
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    if p.runs:
        _set_run_font(p.runs[0], size_pt=size_pt, color=color or THEME.body, bold=bold)


def _add_accent_bar(slide, *, top=Inches(0), height=Inches(0.12)) -> None:
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0),
        top,
        SLIDE_W,
        height,
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = THEME.accent
    shape.line.fill.background()


def _add_footer(slide, deck_title: str) -> None:
    _add_textbox(
        slide,
        Inches(0.55),
        Inches(7.05),
        Inches(5),
        Inches(0.35),
        "AI-AGENT",
        size_pt=9,
        color=THEME.muted,
    )
    if deck_title:
        _add_textbox(
            slide,
            Inches(8),
            Inches(7.05),
            Inches(4.8),
            Inches(0.35),
            deck_title[:60],
            size_pt=9,
            color=THEME.muted,
            align=PP_ALIGN.RIGHT,
        )


def _add_title_slide(prs: Presentation, slide_data: dict[str, Any], deck_title: str) -> None:
    slide = prs.slides.add_slide(_blank_layout(prs))
    _set_slide_bg(slide, THEME.bg_dark)

    # Decorative accent panel
    panel = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(5.85), SLIDE_W, Inches(1.65))
    panel.fill.solid()
    panel.fill.fore_color.rgb = THEME.accent
    panel.line.fill.background()

    title = str(slide_data.get("title") or slide_data.get("heading") or deck_title)
    subtitle = str(slide_data.get("subtitle") or slide_data.get("body") or "")

    _add_textbox(
        slide,
        Inches(0.85),
        Inches(1.6),
        Inches(11.5),
        Inches(1.4),
        title,
        size_pt=40,
        color=THEME.on_dark,
        bold=True,
    )
    if subtitle:
        _add_textbox(
            slide,
            Inches(0.88),
            Inches(3.2),
            Inches(11),
            Inches(1.2),
            subtitle,
            size_pt=20,
            color=THEME.on_dark_sub,
        )

    _add_textbox(
        slide,
        Inches(0.88),
        Inches(6.15),
        Inches(6),
        Inches(0.4),
        "AI-AGENT · 数字员工报告",
        size_pt=11,
        color=THEME.on_dark,
    )


def _add_section_slide(prs: Presentation, title: str, deck_title: str) -> None:
    slide = prs.slides.add_slide(_blank_layout(prs))
    _set_slide_bg(slide, THEME.accent_soft)
    _add_accent_bar(slide, top=Inches(3.35), height=Inches(0.08))
    _add_textbox(
        slide,
        Inches(0.9),
        Inches(2.6),
        Inches(11.5),
        Inches(1.2),
        title,
        size_pt=32,
        color=THEME.title,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    _add_footer(slide, deck_title)


def _add_content_slide(
    prs: Presentation,
    title: str,
    bullets: list[str] | None,
    body: str | None,
    deck_title: str,
) -> None:
    slide = prs.slides.add_slide(_blank_layout(prs))
    _set_slide_bg(slide, THEME.bg_light)
    _add_accent_bar(slide)

    _add_textbox(
        slide,
        Inches(0.65),
        Inches(0.35),
        Inches(12),
        Inches(0.75),
        title,
        size_pt=26,
        color=THEME.title,
        bold=True,
    )

    items = [str(b).strip() for b in (bullets or []) if str(b).strip()]
    if not items and body:
        items = [str(body)]

    if items:
        box = slide.shapes.add_textbox(Inches(0.75), Inches(1.35), Inches(11.8), Inches(5.4))
        tf = box.text_frame
        tf.word_wrap = True
        for i, item in enumerate(items[:12]):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = f"• {item}" if not item.startswith("•") else item
            p.space_after = Pt(10)
            if p.runs:
                _set_run_font(p.runs[0], size_pt=18, color=THEME.body)
    _add_footer(slide, deck_title)


def _add_table_slide(prs: Presentation, title: str, rows: list[list[str]], deck_title: str) -> None:
    if not rows:
        return
    slide = prs.slides.add_slide(_blank_layout(prs))
    _set_slide_bg(slide, THEME.bg_light)
    _add_accent_bar(slide)
    _add_textbox(
        slide,
        Inches(0.65),
        Inches(0.35),
        Inches(12),
        Inches(0.75),
        title,
        size_pt=24,
        color=THEME.title,
        bold=True,
    )

    cols = max(len(r) for r in rows)
    table_shape = slide.shapes.add_table(
        len(rows),
        cols,
        Inches(0.65),
        Inches(1.25),
        Inches(12),
        Inches(min(5.5, 0.45 * len(rows) + 0.5)),
    )
    table = table_shape.table
    for r_idx, row in enumerate(rows):
        for c_idx in range(cols):
            cell = table.cell(r_idx, c_idx)
            text = row[c_idx] if c_idx < len(row) else ""
            cell.text = text
            for paragraph in cell.text_frame.paragraphs:
                for run in paragraph.runs:
                    is_header = r_idx == 0
                    _set_run_font(
                        run,
                        size_pt=12 if is_header else 11,
                        color=THEME.title if is_header else THEME.body,
                        bold=is_header,
                    )
            if r_idx == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = THEME.accent_soft
            elif r_idx % 2 == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    _add_footer(slide, deck_title)


def build_presentation(
    *,
    slides: list[dict[str, Any]],
    deck_title: str = "Presentation",
) -> Presentation:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    prs.core_properties.title = deck_title[:200]
    prs.core_properties.author = "AI-AGENT"

    if not slides:
        _add_title_slide(prs, {"title": deck_title}, deck_title)
        return prs

    for idx, slide_data in enumerate(slides):
        layout = str(slide_data.get("layout") or "").lower()
        title = str(slide_data.get("title") or slide_data.get("heading") or f"Slide {idx + 1}")
        body = slide_data.get("body")
        bullets = slide_data.get("bullets") or slide_data.get("points")
        if bullets and not isinstance(bullets, list):
            bullets = [str(bullets)]
        if body and not bullets:
            bullets = [str(body)]

        table = slide_data.get("table")
        if layout == "table" or table:
            rows = table if isinstance(table, list) else []
            _add_table_slide(prs, title, rows, deck_title)
            continue

        if layout == "section":
            _add_section_slide(prs, title, deck_title)
            continue

        is_title = layout == "title" or (
            idx == 0
            and not bullets
            and (slide_data.get("subtitle") or not slide_data.get("body"))
        )
        if is_title:
            _add_title_slide(prs, slide_data, deck_title)
            continue

        _add_content_slide(prs, title, bullets, str(body) if body else None, deck_title)

    if not prs.slides:
        _add_title_slide(prs, {"title": deck_title}, deck_title)

    return prs


def markdown_to_slides(markdown: str, *, title: str) -> list[dict[str, Any]]:
    """Parse markdown report into structured slide data."""
    from markdown_it import MarkdownIt

    md = MarkdownIt("commonmark", {"breaks": True}).enable("table")
    tokens = md.parse(markdown.strip())
    slides: list[dict[str, Any]] = []
    current: dict[str, Any] | None = None
    i = 0

    def flush() -> None:
        nonlocal current
        if current and (current.get("title") or current.get("bullets") or current.get("table")):
            slides.append(current)
        current = None

    while i < len(tokens):
        tok = tokens[i]
        if tok.type == "heading_open":
            level = int(tok.tag[1])
            text = tokens[i + 1].content.strip()
            i += 3
            if level == 1 and not slides and not current:
                current = {"layout": "title", "title": text, "subtitle": ""}
            elif level == 1:
                flush()
                current = {"layout": "section", "title": text}
                flush()
            else:
                flush()
                current = {"title": text, "bullets": []}
        elif tok.type == "paragraph_open":
            inline = tokens[i + 1]
            text = _inline_plain(inline)
            i += 3
            if current is None:
                current = {"title": title, "bullets": []}
            if current.get("layout") == "title" and not current.get("subtitle"):
                current["subtitle"] = text
            elif text:
                current.setdefault("bullets", []).append(text)
        elif tok.type == "bullet_list_open":
            i += 1
            while i < len(tokens) and tokens[i].type != "bullet_list_close":
                if tokens[i].type == "list_item_open":
                    if tokens[i + 1].type == "paragraph_open":
                        inline = tokens[i + 2]
                        text = _inline_plain(inline)
                        if current is None:
                            current = {"title": title, "bullets": []}
                        current.setdefault("bullets", []).append(text)
                        i += 5
                        continue
                i += 1
            i += 1
        elif tok.type == "table_open":
            table_title = (current or {}).get("title") or title
            flush()
            rows: list[list[str]] = []
            i += 1
            while i < len(tokens) and tokens[i].type != "table_close":
                if tokens[i].type == "tr_open":
                    row: list[str] = []
                    i += 1
                    while i < len(tokens) and tokens[i].type != "tr_close":
                        if tokens[i].type in ("th_open", "td_open"):
                            row.append(_inline_plain(tokens[i + 1]))
                            i += 3
                            continue
                        i += 1
                    if row and not all(re.match(r"^[-:]+$", c) for c in row):
                        rows.append(row)
                i += 1
            flush()
            if rows:
                slides.append({
                    "layout": "table",
                    "title": table_title,
                    "table": rows,
                })
            i += 1
        elif tok.type == "fence":
            code = tok.content.strip()
            if current is None:
                current = {"title": title, "bullets": []}
            current.setdefault("bullets", []).append(code[:800])
            i += 1
        elif tok.type == "hr":
            flush()
            i += 1
        else:
            i += 1

    flush()

    if not slides:
        slides = [{"layout": "title", "title": title, "subtitle": markdown[:200]}]
    elif slides[0].get("layout") != "title":
        slides.insert(0, {"layout": "title", "title": title, "subtitle": ""})

    return slides


def _inline_plain(inline) -> str:
    if not inline.children:
        return (inline.content or "").strip()
    parts: list[str] = []
    for child in inline.children:
        if child.type in ("text", "code_inline"):
            parts.append(child.content)
        elif child.type in ("softbreak", "hardbreak"):
            parts.append(" ")
    return "".join(parts).strip()


def export_pptx_file(
    *,
    output_path: Path,
    slides_json: str,
    deck_title: str = "Presentation",
) -> Path:
    slides = _parse_slides(slides_json)
    prs = build_presentation(slides=slides, deck_title=deck_title)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output_path))
    return output_path


def export_markdown_to_pptx(markdown: str, *, title: str) -> bytes:
    import io

    slides = markdown_to_slides(markdown, title=title)
    prs = build_presentation(slides=slides, deck_title=title)
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
