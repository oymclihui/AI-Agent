"""Export Markdown to styled Word (.docx) matching frontend MarkdownView (light theme)."""
from __future__ import annotations

import io
import re
from dataclasses import dataclass

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from markdown_it import MarkdownIt


@dataclass(frozen=True)
class MarkdownDocxTheme:
    """Aligned with frontend globals.css MarkdownView light theme."""

    heading: RGBColor = RGBColor(0x1E, 0x3A, 0x5F)
    body: RGBColor = RGBColor(0x1E, 0x3A, 0x5F)
    muted: RGBColor = RGBColor(0x5B, 0x7A, 0x9A)
    link: RGBColor = RGBColor(0x25, 0x63, 0xEB)
    accent: RGBColor = RGBColor(0x3B, 0x82, 0xF6)
    border: str = "DBEAFE"
    code_inline_bg: str = "E0F2FE"
    code_inline_text: RGBColor = RGBColor(0x03, 0x69, 0xA1)
    code_block_bg: str = "F6F8FA"
    code_block_text: RGBColor = RGBColor(0x24, 0x29, 0x2F)
    blockquote_bg: str = "EFF6FF"
    table_header_bg: str = "E8F2FF"
    table_zebra_bg: str = "F8FBFF"
    font_latin: str = "Calibri"
    font_cjk: str = "Microsoft YaHei UI"
    body_pt: float = 11.0
    h1_pt: float = 17.0
    h2_pt: float = 14.5
    h3_pt: float = 12.5
    h4_pt: float = 11.0


THEME = MarkdownDocxTheme()
_MD = MarkdownIt("commonmark", {"breaks": True}).enable("table").enable("strikethrough")


def export_markdown_to_docx(markdown: str, *, title: str | None = None) -> bytes:
    doc = Document()
    _configure_document(doc, title=title)
    renderer = _DocxMarkdownRenderer(doc, THEME)
    renderer.render(markdown.strip())
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _configure_document(doc: Document, *, title: str | None) -> None:
    if title:
        doc.core_properties.title = title[:200]

    section = doc.sections[0]
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.85)
    section.left_margin = Inches(0.95)
    section.right_margin = Inches(0.95)

    normal = doc.styles["Normal"]
    normal.font.name = THEME.font_latin
    normal.font.size = Pt(THEME.body_pt)
    normal.font.color.rgb = THEME.body
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    normal.paragraph_format.line_spacing = 1.75
    normal.paragraph_format.space_after = Pt(6)
    _set_style_east_asia(normal, THEME.font_cjk)


def _set_style_east_asia(style, font_name: str) -> None:
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), font_name)


def _set_run_fonts(run, theme: MarkdownDocxTheme) -> None:
    run.font.name = theme.font_latin
    r = run._element
    rpr = r.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), theme.font_cjk)


def _set_paragraph_bottom_border(paragraph, color_hex: str, size: int = 6) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:color"), color_hex)
    p_bdr.append(bottom)
    ppr.append(p_bdr)


def _set_cell_shading(cell, fill_hex: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill_hex.lstrip("#"))
    shading.set(qn("w:val"), "clear")
    cell._tc.get_or_add_tcPr().append(shading)


def _add_hyperlink(paragraph, text: str, url: str, theme: MarkdownDocxTheme) -> None:
    if not url:
        run = paragraph.add_run(text)
        _apply_hyperlink_style(run, theme)
        return
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "2563EB")
    r_pr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), theme.font_latin)
    rfonts.set(qn("w:hAnsi"), theme.font_latin)
    rfonts.set(qn("w:eastAsia"), theme.font_cjk)
    r_pr.append(rfonts)
    new_run.append(r_pr)
    text_el = OxmlElement("w:t")
    text_el.text = text
    new_run.append(text_el)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def _apply_hyperlink_style(run, theme: MarkdownDocxTheme) -> None:
    _set_run_fonts(run, theme)
    run.font.color.rgb = theme.link
    run.font.underline = True


def _set_cell_borders(cell, color_hex: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), color_hex)
        borders.append(el)
    tc_pr.append(borders)


class _DocxMarkdownRenderer:
    def __init__(self, doc: Document, theme: MarkdownDocxTheme) -> None:
        self.doc = doc
        self.theme = theme
        self._list_depth = 0
        self._list_ordered = False
        self._in_list_item = False
        self._in_blockquote = False

    def render(self, markdown: str) -> None:
        tokens = _MD.parse(markdown)
        i = 0
        while i < len(tokens):
            tok = tokens[i]
            if tok.type == "heading_open":
                level = int(tok.tag[1])
                inline = tokens[i + 1]
                i += 3
                self._add_heading(inline, level)
            elif tok.type == "paragraph_open":
                inline = tokens[i + 1]
                i += 3
                if self._in_blockquote:
                    self._add_blockquote_paragraph(inline)
                else:
                    self._add_body_paragraph(inline)
            elif tok.type == "bullet_list_open":
                self._list_depth += 1
                self._list_ordered = False
                i += 1
            elif tok.type == "bullet_list_close":
                self._list_depth = max(0, self._list_depth - 1)
                i += 1
            elif tok.type == "ordered_list_open":
                self._list_depth += 1
                self._list_ordered = True
                i += 1
            elif tok.type == "ordered_list_close":
                self._list_depth = max(0, self._list_depth - 1)
                i += 1
            elif tok.type == "list_item_open":
                self._in_list_item = True
                i += 1
            elif tok.type == "list_item_close":
                self._in_list_item = False
                i += 1
            elif tok.type == "blockquote_open":
                self._in_blockquote = True
                i += 1
            elif tok.type == "blockquote_close":
                self._in_blockquote = False
                i += 1
            elif tok.type == "fence":
                self._add_code_block(tok.content.rstrip("\n"), tok.info or "")
                i += 1
            elif tok.type == "hr":
                self._add_hr()
                i += 1
            elif tok.type == "table_open":
                i, _ = self._add_table(tokens, i)
            else:
                i += 1

    def _add_heading(self, inline, level: int) -> None:
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14 if level <= 2 else 10)
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run(self._inline_text(inline))
        _set_run_fonts(run, self.theme)
        run.bold = True
        run.font.color.rgb = self.theme.heading if level <= 2 else self.theme.body
        sizes = {1: self.theme.h1_pt, 2: self.theme.h2_pt, 3: self.theme.h3_pt, 4: self.theme.h4_pt}
        run.font.size = Pt(sizes.get(level, self.theme.h4_pt))
        if level == 1:
            _set_paragraph_bottom_border(p, self.theme.border)

    def _add_body_paragraph(self, inline) -> None:
        if self._in_list_item:
            style = "List Number" if self._list_ordered else "List Bullet"
            try:
                p = self.doc.add_paragraph(style=style)
            except KeyError:
                p = self.doc.add_paragraph()
            indent = max(0, self._list_depth - 1) * 0.25
            p.paragraph_format.left_indent = Inches(0.12 + indent)
        else:
            p = self.doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        self._render_inline(p, inline)

    def _add_blockquote_paragraph(self, inline) -> None:
        p = self.doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.18)
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        ppr = p._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), self.theme.blockquote_bg)
        shd.set(qn("w:val"), "clear")
        ppr.append(shd)
        p_bdr = OxmlElement("w:pBdr")
        left = OxmlElement("w:left")
        left.set(qn("w:val"), "single")
        left.set(qn("w:sz"), "18")
        left.set(qn("w:color"), "3B82F6")
        p_bdr.append(left)
        ppr.append(p_bdr)
        self._render_inline(p, inline, muted=True)

    def _add_code_block(self, code: str, lang: str) -> None:
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(8)
        ppr = p._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), self.theme.code_block_bg)
        shd.set(qn("w:val"), "clear")
        ppr.append(shd)
        _set_paragraph_bottom_border(p, self.theme.border, size=4)
        run = p.add_run(code)
        _set_run_fonts(run, self.theme)
        run.font.name = "Consolas"
        rpr = run._element.get_or_add_rPr()
        rfonts = rpr.rFonts
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.append(rfonts)
        rfonts.set(qn("w:ascii"), "Consolas")
        rfonts.set(qn("w:hAnsi"), "Consolas")
        run.font.size = Pt(9.5)
        run.font.color.rgb = self.theme.code_block_text

    def _add_hr(self) -> None:
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(12)
        _set_paragraph_bottom_border(p, self.theme.border, size=4)

    def _add_table(self, tokens: list, start: int) -> tuple[int, None]:
        rows: list[list[str]] = []
        header_rows = 0
        i = start
        in_head = False
        while i < len(tokens):
            tok = tokens[i]
            if tok.type == "table_close":
                break
            if tok.type == "thead_open":
                in_head = True
            elif tok.type == "thead_close":
                in_head = False
            elif tok.type == "tr_open":
                row: list[str] = []
                i += 1
                while i < len(tokens) and tokens[i].type != "tr_close":
                    if tokens[i].type in ("th_open", "td_open"):
                        inline = tokens[i + 1]
                        row.append(self._inline_text(inline))
                        i += 3
                        continue
                    i += 1
                rows.append(row)
                if in_head and header_rows == 0:
                    header_rows = 1
            i += 1

        if not rows:
            return i + 1, None

        cols = max(len(r) for r in rows)
        table = self.doc.add_table(rows=len(rows), cols=cols)
        table.style = "Table Grid"
        table.autofit = True

        for r_idx, row in enumerate(rows):
            for c_idx in range(cols):
                cell = table.rows[r_idx].cells[c_idx]
                text = row[c_idx] if c_idx < len(row) else ""
                cell.text = ""
                p = cell.paragraphs[0]
                run = p.add_run(text)
                _set_run_fonts(run, self.theme)
                run.font.size = Pt(10)
                if r_idx < header_rows:
                    run.bold = True
                    run.font.color.rgb = self.theme.heading
                    _set_cell_shading(cell, self.theme.table_header_bg)
                elif r_idx % 2 == 0:
                    _set_cell_shading(cell, self.theme.table_zebra_bg)
                _set_cell_borders(cell, self.theme.border)

        gap = self.doc.add_paragraph()
        gap.paragraph_format.space_after = Pt(4)
        return i + 1, None

    def _render_inline(self, paragraph, inline, *, muted: bool = False) -> None:
        if not inline.children:
            text = inline.content or ""
            if text:
                run = paragraph.add_run(text)
                self._apply_run_style(run, muted=muted)
            return

        style_stack: list[str] = []
        link_href: str | None = None
        link_text_parts: list[str] = []
        i = 0
        children = inline.children

        def flush_link() -> None:
            nonlocal link_text_parts, link_href
            if link_href and link_text_parts:
                _add_hyperlink(paragraph, "".join(link_text_parts), link_href, self.theme)
            link_text_parts = []
            link_href = None

        while i < len(children):
            child = children[i]
            if child.type == "text":
                if "link" in style_stack:
                    link_text_parts.append(child.content)
                else:
                    run = paragraph.add_run(child.content)
                    self._apply_run_style(run, style_stack, muted=muted)
            elif child.type == "code_inline":
                flush_link()
                run = paragraph.add_run(child.content)
                _set_run_fonts(run, self.theme)
                run.font.size = Pt(9.5)
                run.font.color.rgb = self.theme.code_inline_text
                run.font.name = "Consolas"
                rpr = run._element.get_or_add_rPr()
                shd = OxmlElement("w:shd")
                shd.set(qn("w:fill"), self.theme.code_inline_bg)
                shd.set(qn("w:val"), "clear")
                rpr.append(shd)
            elif child.type == "softbreak":
                flush_link()
                paragraph.add_run("\n")
            elif child.type == "hardbreak":
                flush_link()
                paragraph.add_run("\n")
            elif child.type == "strong_open":
                style_stack.append("strong")
            elif child.type == "strong_close":
                if style_stack and style_stack[-1] == "strong":
                    style_stack.pop()
            elif child.type == "em_open":
                style_stack.append("em")
            elif child.type == "em_close":
                if style_stack and style_stack[-1] == "em":
                    style_stack.pop()
            elif child.type == "link_open":
                flush_link()
                link_href = child.attrGet("href")
                style_stack.append("link")
            elif child.type == "link_close":
                flush_link()
                if style_stack and style_stack[-1] == "link":
                    style_stack.pop()
            elif child.type == "s_open":
                style_stack.append("strike")
            elif child.type == "s_close":
                if style_stack and style_stack[-1] == "strike":
                    style_stack.pop()
            i += 1
        flush_link()

    def _apply_run_style(
        self,
        run,
        style_stack: list[str] | None = None,
        *,
        muted: bool = False,
    ) -> None:
        stack = style_stack or []
        _set_run_fonts(run, self.theme)
        run.font.size = Pt(self.theme.body_pt)
        run.font.color.rgb = self.theme.muted if muted else self.theme.body
        if "strong" in stack:
            run.bold = True
            run.font.color.rgb = self.theme.heading
        if "em" in stack:
            run.italic = True
        if "strike" in stack:
            run.font.strike = True
        if "link" in stack:
            run.font.color.rgb = self.theme.link
            run.font.underline = True

    @staticmethod
    def _inline_text(inline) -> str:
        if not inline.children:
            return inline.content or ""
        parts: list[str] = []
        for child in inline.children:
            if child.type == "text":
                parts.append(child.content)
            elif child.type == "code_inline":
                parts.append(child.content)
            elif child.type in ("softbreak", "hardbreak"):
                parts.append("\n")
        return "".join(parts)


def strip_yaml_front_matter(markdown: str) -> str:
    if markdown.startswith("---"):
        m = re.match(r"^---\n.*?\n---\n", markdown, re.DOTALL)
        if m:
            return markdown[m.end() :]
    return markdown
