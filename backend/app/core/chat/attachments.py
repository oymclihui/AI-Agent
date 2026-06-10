"""Parse chat attachments into LLM-readable text."""
from __future__ import annotations

import json
import os
import tempfile
from typing import Any

import structlog

from app.knowledge.loader import DocumentLoader

logger = structlog.get_logger(__name__)

# No upload byte cap — only truncate extracted text for LLM context window.
MAX_EXTRACT_CHARS_PER_FILE = 200_000
MAX_EXTRACT_CHARS_TOTAL = 600_000

SUPPORTED_EXTENSIONS = {
    "pdf", "docx", "doc", "xlsx", "xls", "pptx", "ppt",
    "txt", "md", "rst", "csv", "json", "xml", "html", "htm",
    "jpg", "jpeg", "png", "gif", "webp", "bmp",
}


def _truncate(text: str, limit: int) -> tuple[str, bool]:
    if len(text) <= limit:
        return text, False
    return text[:limit] + f"\n\n…（内容已截断，原文约 {len(text):,} 字符）", True


async def _load_xlsx(file_path: str, filename: str) -> str:
    import pandas as pd

    parts: list[str] = []
    xl = pd.ExcelFile(file_path)
    for sheet in xl.sheet_names:
        df = pd.read_excel(file_path, sheet_name=sheet)
        preview_rows = min(len(df), 200)
        preview = df.head(preview_rows)
        parts.append(
            f"### 工作表: {sheet}\n"
            f"行数={len(df)}, 列={list(df.columns)}\n"
            f"{preview.to_string(index=False)}"
        )
        if len(df) > preview_rows:
            parts.append(f"…（仅展示前 {preview_rows} 行）")
    return "\n\n".join(parts)


async def _load_pptx(file_path: str, filename: str) -> str:
    from pptx import Presentation

    prs = Presentation(file_path)
    parts: list[str] = []
    for idx, slide in enumerate(prs.slides, start=1):
        texts: list[str] = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text and shape.text.strip():
                texts.append(shape.text.strip())
        if texts:
            parts.append(f"### 第 {idx} 页\n" + "\n".join(texts))
        else:
            parts.append(f"### 第 {idx} 页\n（无提取到文本）")
    return "\n\n".join(parts) if parts else "（PPT 中未提取到文本）"


async def _load_image(file_path: str, filename: str) -> str:
    from PIL import Image

    with Image.open(file_path) as img:
        w, h = img.size
        mode = img.mode
        fmt = img.format or "unknown"
    return (
        f"[图片文件: {filename}]\n"
        f"格式={fmt}, 尺寸={w}x{h}, 色彩模式={mode}\n"
        "说明：当前为文本分析模式，无法直接识别图片像素内容。"
        "请结合用户问题，说明需要 OCR/视觉模型时可另行处理。"
    )


async def _load_csv(file_path: str, filename: str) -> str:
    import pandas as pd

    df = pd.read_csv(file_path)
    preview_rows = min(len(df), 200)
    body = df.head(preview_rows).to_string(index=False)
    suffix = f"\n…（仅展示前 {preview_rows} 行，共 {len(df)} 行）" if len(df) > preview_rows else ""
    return f"列={list(df.columns)}\n{body}{suffix}"


async def extract_attachment_text(file_path: str, filename: str) -> dict[str, Any]:
    suffix = filename.lower().rsplit(".", 1)[-1] if "." in filename else "txt"
    if suffix not in SUPPORTED_EXTENSIONS:
        suffix = "txt"

    try:
        if suffix in ("xlsx", "xls"):
            raw = await _load_xlsx(file_path, filename)
            kind = "excel"
        elif suffix in ("pptx", "ppt"):
            raw = await _load_pptx(file_path, filename)
            kind = "pptx"
        elif suffix == "csv":
            raw = await _load_csv(file_path, filename)
            kind = "csv"
        elif suffix == "json":
            with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                data = json.load(f)
            raw = json.dumps(data, ensure_ascii=False, indent=2)
            kind = "json"
        elif suffix == "xml":
            with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                raw = f.read()
            kind = "xml"
        elif suffix in ("jpg", "jpeg", "png", "gif", "webp", "bmp"):
            raw = await _load_image(file_path, filename)
            kind = "image"
        else:
            docs = await DocumentLoader.load_file(file_path, filename)
            raw = "\n\n".join(d.page_content for d in docs if d.page_content.strip())
            kind = suffix

        text, truncated = _truncate(raw or "（未能提取文本内容）", MAX_EXTRACT_CHARS_PER_FILE)
        return {
            "filename": filename,
            "kind": kind,
            "text": text,
            "char_count": len(raw or ""),
            "truncated": truncated,
            "ok": True,
        }
    except Exception as exc:
        logger.warning("attachment_extract_failed", filename=filename, error=str(exc))
        return {
            "filename": filename,
            "kind": suffix,
            "text": f"（解析失败: {exc}）",
            "char_count": 0,
            "truncated": False,
            "ok": False,
            "error": str(exc),
        }


async def build_attachment_context(
    files: list[tuple[str, bytes]],
) -> tuple[str, list[dict[str, Any]]]:
    """Parse multiple uploads; return LLM context block + metadata list."""
    if not files:
        return "", []

    sections: list[str] = []
    meta: list[dict[str, Any]] = []
    total_chars = 0

    for filename, data in files:
        safe_name = os.path.basename(filename) or "attachment"
        with tempfile.NamedTemporaryFile(suffix=f"_{safe_name}", delete=False) as tmp:
            tmp.write(data)
            tmp_path = tmp.name
        try:
            parsed = await extract_attachment_text(tmp_path, safe_name)
        finally:
            os.unlink(tmp_path)

        meta.append({
            "filename": safe_name,
            "size": len(data),
            "kind": parsed.get("kind"),
            "char_count": parsed.get("char_count"),
            "truncated": parsed.get("truncated", False),
            "ok": parsed.get("ok", True),
        })

        block = (
            f"## 附件: {safe_name}\n"
            f"类型: {parsed.get('kind')} | 大小: {len(data):,} 字节\n\n"
            f"{parsed.get('text', '')}"
        )
        remaining = MAX_EXTRACT_CHARS_TOTAL - total_chars
        if remaining <= 0:
            sections.append(f"## 附件: {safe_name}\n（总上下文已达上限，未纳入分析）")
            continue
        if len(block) > remaining:
            block, _ = _truncate(block, remaining)
        sections.append(block)
        total_chars += len(block)

    context = (
        "用户上传了以下附件，请结合用户问题进行分析、摘要或回答。\n"
        "若数据为表格/文档，请指出关键发现、异常与可执行建议。\n\n"
        + "\n\n---\n\n".join(sections)
    )
    return context, meta
