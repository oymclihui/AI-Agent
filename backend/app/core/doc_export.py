"""Convert markdown / plain-text reports to Office formats for download."""
from __future__ import annotations

import io
import re
from pathlib import Path
from typing import Literal

DownloadFormat = Literal["md", "docx", "xlsx", "pptx"]

DOWNLOAD_FORMATS: tuple[DownloadFormat, ...] = ("md", "docx", "xlsx", "pptx")

FORMAT_MEDIA: dict[str, str] = {
    "md": "text/markdown; charset=utf-8",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
}


def normalize_format(value: str | None, *, default: str = "md") -> DownloadFormat:
    fmt = (value or default).lower().strip().lstrip(".")
    if fmt not in DOWNLOAD_FORMATS:
        raise ValueError(f"Unsupported format: {value!r} (use md, docx, xlsx, or pptx)")
    return fmt  # type: ignore[return-value]


def formats_for_filename(filename: str) -> list[DownloadFormat]:
    ext = Path(filename).suffix.lower().lstrip(".")
    if ext in ("pptx", "xlsx", "docx"):
        return [ext]  # type: ignore[list-item]
    return list(DOWNLOAD_FORMATS)


def replace_extension(filename: str, fmt: str) -> str:
    stem = Path(filename).stem or "report"
    return f"{stem}.{fmt}"


def convert_markdown(markdown: str, fmt: DownloadFormat, *, title: str = "Report") -> bytes:
    if fmt == "md":
        return markdown.encode("utf-8")
    if fmt == "docx":
        return _markdown_to_docx(markdown, title=title)
    if fmt == "xlsx":
        return _markdown_to_xlsx(markdown, title=title)
    return _markdown_to_pptx(markdown, title=title)


def _markdown_to_docx(markdown: str, *, title: str) -> bytes:
    from app.core.design.docx_markdown_export import export_markdown_to_docx, strip_yaml_front_matter

    return export_markdown_to_docx(strip_yaml_front_matter(markdown), title=title)


def _parse_markdown_tables(markdown: str) -> list[list[list[str]]]:
    tables: list[list[list[str]]] = []
    lines = markdown.splitlines()
    i = 0
    while i < len(lines):
        if not lines[i].strip().startswith("|"):
            i += 1
            continue
        block: list[list[str]] = []
        while i < len(lines) and lines[i].strip().startswith("|"):
            row = lines[i].strip()
            if re.match(r"^\|[\s\-:|]+\|$", row):
                i += 1
                continue
            cells = [c.strip() for c in row.strip("|").split("|")]
            block.append(cells)
            i += 1
        if block:
            tables.append(block)
    return tables


def _markdown_to_xlsx(markdown: str, *, title: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    summary = wb.active
    summary.title = "报告"
    summary["A1"] = title
    summary["A1"].font = Font(bold=True, size=14)
    row = 3
    for line in markdown.splitlines():
        text = line.rstrip()
        if not text or (text.startswith("|") and "|" in text[1:]):
            continue
        if text.startswith("#"):
            text = text.lstrip("#").strip()
            summary.cell(row=row, column=1, value=text).font = Font(bold=True)
        else:
            summary.cell(row=row, column=1, value=text)
        row += 1

    tables = _parse_markdown_tables(markdown)
    for idx, table in enumerate(tables, start=1):
        ws = wb.create_sheet(title=f"表{idx}"[:31])
        for r, cells in enumerate(table, start=1):
            for c, val in enumerate(cells, start=1):
                ws.cell(row=r, column=c, value=val)
        if table:
            for cell in ws[1]:
                cell.font = Font(bold=True)

    if not tables and row <= 3:
        summary["A3"] = markdown.strip()

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _markdown_to_pptx(markdown: str, *, title: str) -> bytes:
    from app.core.design.pptx_export import export_markdown_to_pptx

    return export_markdown_to_pptx(markdown, title=title)
