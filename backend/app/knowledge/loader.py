"""
Document Loader - supports PDF, DOCX, plain text, Excel, CSV, and URLs.
Inspired by LangChain's document loaders and Dify's ingestion pipeline.
"""
from datetime import date, datetime
from typing import Optional
import structlog
from langchain_core.documents import Document

logger = structlog.get_logger(__name__)

_EXCEL_SUFFIXES = {"xlsx", "xlsm", "xls"}
_TEXT_SUFFIXES = {"txt", "md", "rst", "csv", "tsv"}
_MAX_EXCEL_ROWS = 5000


class DocumentLoader:
    @staticmethod
    def _detect_encoding(raw: bytes) -> str:
        """Auto-detect file encoding, prioritize UTF-8 then Chinese encodings."""
        # BOM check
        if raw.startswith(b"\xef\xbb\xbf"):
            return "utf-8-sig"
        if raw.startswith(b"\xff\xfe") or raw.startswith(b"\xfe\xff"):
            return "utf-16"
        # try UTF-8 strict first
        try:
            raw.decode("utf-8")
            return "utf-8"
        except UnicodeDecodeError:
            pass
        # try Chinese encodings (gb18030 is superset of gbk/gb2312)
        for enc in ("gb18030", "gbk", "gb2312", "big5"):
            try:
                raw.decode(enc)
                return enc
            except (UnicodeDecodeError, LookupError):
                continue
        # chardet as last resort
        try:
            import chardet
            result = chardet.detect(raw[:8192])
            if result and result.get("encoding"):
                return result["encoding"]
        except ImportError:
            pass
        return "utf-8"

    @staticmethod
    async def load_file(file_path: str, filename: str) -> list[Document]:
        suffix = filename.lower().rsplit(".", 1)[-1] if "." in filename else "txt"
        logger.info("document_load", filename=filename, type=suffix)

        if suffix == "pdf":
            return await DocumentLoader._load_pdf(file_path, filename)
        elif suffix in ("docx", "doc"):
            return await DocumentLoader._load_docx(file_path, filename)
        elif suffix in _EXCEL_SUFFIXES:
            return await DocumentLoader._load_excel(file_path, filename)
        elif suffix in ("html", "htm"):
            return await DocumentLoader._load_html(file_path, filename)
        elif suffix in _TEXT_SUFFIXES:
            return await DocumentLoader._load_text(file_path, filename)
        elif suffix in {"png", "jpg", "jpeg", "webp", "bmp", "gif"}:
            return await DocumentLoader._load_image(file_path, filename)
        else:
            return await DocumentLoader._load_text(file_path, filename)

    @staticmethod
    async def _load_image(file_path: str, filename: str) -> list[Document]:
        from app.knowledge.ocr import image_file_to_text

        text = image_file_to_text(file_path)
        return [Document(page_content=text, metadata={"source": filename, "type": "image"})]

    @staticmethod
    async def load_url(url: str) -> list[Document]:
        import httpx
        from bs4 import BeautifulSoup

        async with httpx.AsyncClient(timeout=20.0) as client:
            response = await client.get(url, follow_redirects=True)

        soup = BeautifulSoup(response.text, "html.parser")
        for tag in soup(["script", "style", "nav", "footer", "header"]):
            tag.decompose()

        text = soup.get_text(separator="\n", strip=True)
        return [Document(page_content=text, metadata={"source": url, "type": "url"})]

    @staticmethod
    async def _load_pdf(file_path: str, filename: str) -> list[Document]:
        import pypdf
        docs = []
        with open(file_path, "rb") as f:
            reader = pypdf.PdfReader(f)
            for i, page in enumerate(reader.pages):
                text = page.extract_text() or ""
                if text.strip():
                    docs.append(Document(
                        page_content=text,
                        metadata={"source": filename, "page": i + 1, "type": "pdf"},
                    ))
        return docs

    @staticmethod
    async def _load_docx(file_path: str, filename: str) -> list[Document]:
        from docx import Document as DocxDoc
        doc = DocxDoc(file_path)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        full_text = "\n".join(paragraphs)
        return [Document(
            page_content=full_text,
            metadata={"source": filename, "type": "docx"},
        )]

    @staticmethod
    def _cell_text(value) -> str:
        if value is None:
            return ""
        if isinstance(value, bool):
            return "TRUE" if value else "FALSE"
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
        if isinstance(value, datetime):
            return value.isoformat(sep=" ", timespec="seconds")
        if isinstance(value, date):
            return value.isoformat()
        return str(value).replace("\r", " ").replace("\n", " ").replace("|", "／").strip()

    @staticmethod
    def _rows_to_text(sheet_name: str, rows) -> str:
        cleaned: list[list[str]] = []
        truncated = False
        for row in rows:
            if len(cleaned) >= _MAX_EXCEL_ROWS + 1:
                truncated = True
                break
            cells = [DocumentLoader._cell_text(c) for c in row]
            while cells and cells[-1] == "":
                cells.pop()
            if any(cells):
                cleaned.append(cells)
        if not cleaned:
            return ""
        header = cleaned[0]
        lines = [f"# 工作表: {sheet_name}"]
        if len(cleaned) == 1:
            lines.append(" | ".join(header))
        else:
            for row in cleaned[1:]:
                parts = []
                for i, val in enumerate(row):
                    if not val:
                        continue
                    key = header[i] if i < len(header) and header[i] else f"列{i + 1}"
                    parts.append(f"{key}={val}")
                if parts:
                    lines.append(f"[{sheet_name}] " + "；".join(parts))
        if truncated:
            lines.append(f"[{sheet_name}] 仅收录前 {_MAX_EXCEL_ROWS} 行")
        return "\n".join(lines)

    @staticmethod
    async def _load_excel(file_path: str, filename: str) -> list[Document]:
        suffix = filename.lower().rsplit(".", 1)[-1]
        sheets: list[tuple[str, list]] = []
        if suffix == "xls":
            import xlrd
            book = xlrd.open_workbook(file_path)
            for sheet in book.sheets():
                rows = []
                for r in range(sheet.nrows):
                    row = []
                    for c in range(sheet.ncols):
                        cell = sheet.cell(r, c)
                        if cell.ctype == xlrd.XL_CELL_DATE:
                            row.append(xlrd.xldate_as_datetime(cell.value, book.datemode))
                        elif cell.ctype == xlrd.XL_CELL_EMPTY:
                            row.append(None)
                        else:
                            row.append(cell.value)
                    rows.append(row)
                sheets.append((sheet.name, rows))
        else:
            from openpyxl import load_workbook
            wb = load_workbook(file_path, read_only=True, data_only=True)
            try:
                for ws in wb.worksheets:
                    sheets.append((ws.title, list(ws.iter_rows(values_only=True))))
            finally:
                wb.close()

        docs = []
        for title, rows in sheets:
            text = DocumentLoader._rows_to_text(title, rows)
            if text.strip():
                docs.append(Document(
                    page_content=text,
                    metadata={"source": filename, "sheet": title, "type": "excel"},
                ))
        logger.info("excel_loaded", filename=filename, sheets=len(docs))
        return docs

    @staticmethod
    async def _load_text(file_path: str, filename: str) -> list[Document]:
        with open(file_path, "rb") as f:
            raw = f.read()
        encoding = DocumentLoader._detect_encoding(raw)
        content = raw.decode(encoding, errors="replace")
        logger.info("text_encoding_detected", filename=filename, encoding=encoding, content_len=len(content))
        return [Document(
            page_content=content,
            metadata={"source": filename, "type": "text"},
        )]

    @staticmethod
    async def _load_html(file_path: str, filename: str) -> list[Document]:
        from bs4 import BeautifulSoup
        with open(file_path, "rb") as f:
            raw = f.read()
        encoding = DocumentLoader._detect_encoding(raw)
        soup = BeautifulSoup(raw.decode(encoding, errors="replace"), "html.parser")
        text = soup.get_text(separator="\n", strip=True)
        return [Document(
            page_content=text,
            metadata={"source": filename, "type": "html"},
        )]
